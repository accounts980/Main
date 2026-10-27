"""Write ISKO_ERP_Sheet_Column_Decision_Register.xlsx - a planning register only
(not a replacement live ERP). Light pastel styling, no freeze panes, no sheet
protection."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_md as BM
import data_misc as DM
from data_columns import BP, AUDIT

HDR = PatternFill("solid", fgColor="DCEAF7")
SEC = PatternFill("solid", fgColor="EAE4F8")
DEC = {"Keep": "E2F0E8", "Simplify": "D9F0EF", "Combine": "FCECC9", "Convert to report": "FFF4D6", "Removed from Phase 1": "F8DCDD",
       "Remove from Phase 1": "F8DCDD", "Remove": "F8DCDD", "Rename": "D9F0EF", "Move to report": "FFF4D6", "Add": "EAE4F8"}
THIN = Side(style="thin", color="D9E0E8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TXT = "243247"


def sheet(wb, name, title, note, headers, rows, widths, dec_col=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=13, color=TXT)
    ws["A2"] = note; ws["A2"].font = Font(italic=True, size=9, color="5A6B82")
    r0 = 4
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h); c.fill = HDR; c.font = Font(bold=True, color=TXT, size=10); c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, str):
                v = v.replace("**", "")
            c = ws.cell(i, j, v); c.border = BORDER; c.font = Font(color=TXT, size=9.5)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if dec_col is not None:
            v = str(row[dec_col]).replace("**", "")
            if v in DEC:
                ws.cell(i, dec_col + 1).fill = PatternFill("solid", fgColor=DEC[v])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + max(1, len(rows))}"
    return ws


def main():
    wb = Workbook(); ws = wb.active; ws.title = "README"; ws.sheet_view.showGridLines = False
    lines = [("ISKO Excel ERP - Sheet & Column Decision Register (v1.0, 27-Oct-2026)", True),
             ("Planning register supporting ISKO_ERP_100K_Row_Resilience_Plan.pdf. It is NOT a replacement live ERP and contains no transaction data.", False),
             ("Measured figures come from the four attached 2026 workbooks (XML audit). Performance targets are NOT yet tested in Excel.", False),
             ("", False),
             ("Sheets in this register:", True),
             ("Sheet_Register - all 120 current sheets with measured size and decision", False),
             ("Column_Decisions - every column of every current Excel Table (incl. MIS imports) with action and final column", False),
             ("Target_Blueprints - final column order, type, class, validation for each core target table", False),
             ("Source_of_Truth, Snapshot_Policy, Controls - ownership and integrity rules", False),
             ("Perf_Risks, Risk_Register - performance risk evidence and project risks", False),
             ("KPIs, Exceptions, Thresholds, Test_Checklist - MIS and acceptance plan (Not Yet Tested)", False),
             ("Decisions, Requirements - management decisions and business requirement decisions", False),
             ("", False),
             ("Scope overrides applied: no purchase requisition; no print; no mail-sending; no physical-count or count-reconciliation features.", True)]
    for i, (t, b) in enumerate(lines, 1):
        ws.cell(i, 1, t).font = Font(bold=b, size=12 if i == 1 else 10, color=TXT)
    ws.column_dimensions["A"].width = 130

    rows = []
    for b in BM.BOOKS:
        for r in BM.sheet_rows(b):
            rows.append([BM.BOOKNAME[b]] + r)
    sheet(wb, "Sheet_Register", "Sheet-by-sheet rebuild decision register", "Size = used range; table data rows; formula cells (measured).",
          ["Workbook", "Current sheet", "Apparent purpose", "Used range; rows; formulas", "Decision", "Reason and control impact", "Proposed final sheet / table",
           "Target rows", "Inputs / outputs", "Dependencies", "Main risks / controls"], rows, [16, 24, 40, 22, 16, 50, 30, 16, 30, 30, 36], dec_col=4)

    rows = []
    src_of = {"qRMReceipt": "tblRMReceipt", "qRMIssue": "tblRMIssue", "qConsReceipt": "tblConsReceipt", "qConsIssue": "tblConsIssue", "qAdjust": "tblAdjust",
              "qCount": "tblCount", "qPOLine": "tblPOLine", "qProduction": "tblProduction", "qMachineLog": "tblMachineLog", "qRejection": "tblRejection",
              "qWIPCount": "tblWIPCount", "qDCHead": "tblDCHead", "qDCLine": "tblDCLine", "qJWCHead": "tblJWCHead", "qJWCLine": "tblJWCLine",
              "qRequirement": "tblRequirement", "qFGCount": "tblFGCount", "qCoaterReturn": "tblCoaterReturn", "qCoaterOpening": "tblCoaterOpening",
              "qRequisition": "tblRequisition", "qCustomer": "tblCustomer", "qVendor": "tblVendor", "qMaterial": "tblMaterial", "qConsumable": "tblConsumable",
              "qBOM": "tblBOM", "qBOMModel": "tblBOMModel", "qMachine": "tblMachine"}
    for t in BM.tables:
        for i, c in enumerate(t["cols"], 1):
            if t["name"].startswith("q"):
                src = src_of.get(t["name"])
                if src:
                    a, tgt, note = BM.col_action("Dispatch" if src in ("tblCustomer", "tblVendor") else "Stores", src, c)
                    note = f"Import mirror of {src}; X_ extract carries only needed columns. " + note
                else:
                    a, tgt, note = ("Combine", "X_ extract / Power Query output", "Power Query output; see Sheet_Register.")
                if t["name"] in ("qCount", "qWIPCount", "qFGCount", "qRequisition", "qMachineLog"):
                    a, tgt = "Remove", "-"
            else:
                a, tgt, note = BM.col_action(t["wb"], t["name"], c)
            rows.append([BM.BOOKNAME[t["wb"]], t["sheet"], t["name"], i, c, "Yes" if c in t["calc"] else "", a, tgt, note])
    sheet(wb, "Column_Decisions", "Current column inventory and actions", "Every column of every Excel Table in the four books. Calc = table calculated column today.",
          ["Workbook", "Sheet", "Table", "#", "Current column", "Calc", "Action", "Final column", "Note"], rows, [16, 24, 18, 5, 26, 6, 14, 34, 50], dec_col=6)

    rows = []
    for k, v in BP.items():
        cols = v["cols"] + (AUDIT if v["audit"] else [])
        for i, c in enumerate(cols, 1):
            rows.append([v["book"], k, v["title"], i] + list(c))
    sheet(wb, "Target_Blueprints", "Target column blueprints", "Class: IN input, DD dropdown, SN snapshot, CE calculated at entry, LK lookup, CF control flag, IM imported, RP report-only. Validation: H hard stop, W warning.",
          ["Book", "Table", "Title", "#", "Final column", "Type / unit", "Class", "M/O", "Source or rule", "Validation", "Why needed"], rows,
          [12, 16, 30, 5, 28, 18, 7, 6, 50, 34, 30])
    sheet(wb, "Source_of_Truth", "Source-of-truth matrix", "One row per data entity.",
          ["Entity", "System of record", "Entered / imported", "Primary key", "Foreign keys", "Maintained by", "Consumed by", "Type"], DM.ENTITY, [26, 20, 22, 24, 30, 20, 30, 18])
    sheet(wb, "Snapshot_Policy", "Snapshot versus live lookup", "", ["Value", "Captured on", "Snapshot or live", "Reason"], DM.SNAPSHOT, [40, 34, 34, 50])
    sheet(wb, "Controls", "Transaction integrity controls", "", ["Control", "Hard stop / warning", "Where", "Mechanism"], DM.CONTROLS, [36, 26, 30, 70])
    sheet(wb, "Perf_Risks", "Performance risk register (measured evidence)", "", ["Risk", "Evidence (measured)", "Applies to", "Design control"], DM.PERF, [30, 70, 26, 60])
    sheet(wb, "Risk_Register", "Project risk register", "", ["ID", "Risk", "Likelihood", "Impact", "Preventive control", "Detection", "Owner", "Fallback"], DM.RISKS, [7, 40, 11, 9, 44, 30, 18, 34])
    sheet(wb, "KPIs", "MIS KPI definitions", "", ["KPI", "Definition", "Source", "Filters", "Owner"], DM.KPIS, [34, 50, 26, 24, 20])
    sheet(wb, "Exceptions", "MIS exception catalogue", "", ["Code", "Exception", "Severity", "Owner"], DM.CHECKS, [10, 70, 12, 22])
    sheet(wb, "Thresholds", "Provisional acceptance thresholds - NOT YET TESTED", "Targets requiring validation on the company's PCs.",
          ["Operation", "10k-25k", "100k / 400k system", "How measured"], DM.THRESH, [44, 18, 22, 60])
    sheet(wb, "Test_Checklist", "Stress-test checklist - NOT YET TESTED", "Record PC spec, Excel version/bitness, calc mode, file size, rows, median time, errors, pass/fail.",
          ["ID", "Operation", "Volumes", "Steps", "Pass condition", "10k", "25k", "50k", "100k", "400k+", "Result"],
          [list(t) + ["", "", "", "", "", "Not yet tested"] for t in DM.TESTS], [7, 30, 30, 50, 34, 9, 9, 9, 9, 9, 16])
    sheet(wb, "Decisions", "Management decisions required", "", ["ID", "Decision", "Why", "Area", "Answer", "Decided by / date"],
          [list(d) + ["", ""] for d in DM.DECISIONS], [7, 70, 44, 12, 30, 20])
    sheet(wb, "Requirements", "Business requirements - decisions", "", ["Area", "Requirement", "Decision", "How"], DM.REQS, [22, 70, 20, 44], dec_col=2)
    out = os.path.join(ROOT, "ISKO_ERP_Sheet_Column_Decision_Register.xlsx")
    wb.save(out); print(out)


if __name__ == "__main__":
    main()
